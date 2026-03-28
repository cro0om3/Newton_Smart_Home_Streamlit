import re
from pathlib import Path
from typing import Set, Tuple, Dict
import zipfile

try:
    from docx import Document as _Docx
except Exception:
    _Docx = None


def _extract_placeholders_from_text(text: str) -> Set[str]:
    """Extract simple top-level placeholders from text like {{ name }}.

    This intentionally ignores dotted expressions (e.g. {{ item.description }})
    to avoid false-positives for loop variables.
    """
    found = set()
    for m in re.findall(r"\{\{\s*([^}]+?)\s*\}\}", text):
        token = m.split('|', 1)[0].strip()  # remove filters
        # Skip dotted/property access or function calls to reduce false-positives
        if '.' in token or '(' in token or ')' in token:
            continue
        # Only accept simple identifiers
        if re.match(r'^[A-Za-z_][A-Za-z0-9_]*$', token):
            found.add(token)
    return found


def _placeholders_from_html(path: Path) -> Set[str]:
    txt = path.read_text(encoding='utf-8')
    return _extract_placeholders_from_text(txt)


def _placeholders_from_docx(path: Path) -> Set[str]:
    # Try python-docx first for robust extraction
    if _Docx is not None:
        try:
            doc = _Docx(str(path))
            chunks = []
            for p in doc.paragraphs:
                chunks.append(p.text)
            for table in doc.tables:
                for r in table.rows:
                    for c in r.cells:
                        chunks.append(c.text)
            raw = "\n".join(chunks)
            return _extract_placeholders_from_text(raw)
        except Exception:
            pass
    # Fallback: read document.xml from docx package and strip tags
    try:
        with zipfile.ZipFile(str(path)) as z:
            with z.open('word/document.xml') as f:
                xml = f.read().decode('utf-8')
                text = re.sub(r'<[^>]+>', ' ', xml)
                return _extract_placeholders_from_text(text)
    except Exception:
        return set()


def _normalize_replacement_key(k: str) -> str:
    k = str(k).strip()
    # remove surrounding braces if present
    k = re.sub(r'^\{+\s*', '', k)
    k = re.sub(r'\s*\}+\s*$', '', k)
    # drop filters/properties (we only validate top-level keys)
    k = k.split('|', 1)[0].split('.', 1)[0].strip()
    return k


def validate_template(template_path: str | Path, replacements: Dict) -> Tuple[Set[str], Set[str]]:
    """Validate placeholders in a template file against provided replacement keys.

    Returns: (missing_keys, extra_keys) where keys are top-level identifiers.
    """
    p = Path(template_path)
    if not p.exists():
        raise FileNotFoundError(p)
    ext = p.suffix.lower()
    if ext in ('.html', '.htm'):
        placeholders = _placeholders_from_html(p)
    elif ext in ('.docx',):
        placeholders = _placeholders_from_docx(p)
    else:
        # Treat as plain text
        placeholders = _extract_placeholders_from_text(p.read_text(encoding='utf-8'))

    repl_keys = set()
    for k in (replacements.keys() if isinstance(replacements, dict) else []):
        nk = _normalize_replacement_key(k)
        if nk:
            repl_keys.add(nk)

    missing = placeholders - repl_keys
    extra = repl_keys - placeholders
    return missing, extra


def format_mismatch_message(missing: Set[str], extra: Set[str]) -> str:
    parts = []
    if missing:
        parts.append(f"Missing keys: {sorted(list(missing))}")
    if extra:
        parts.append(f"Extra keys: {sorted(list(extra))}")
    return ' ; '.join(parts)
